import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()
section = doc.sections[0]
footer = section.footer
p = footer.paragraphs[0]
p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run('Trang ')
fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
p._p.append(fldSimple)

table = doc.add_table(rows=2, cols=2)
tblPr = table._element.xpath('w:tblPr')
borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:left w:val="none"/><w:right w:val="none"/><w:insideV w:val="none"/></w:tblBorders>')
tblPr[0].append(borders)

cell = table.cell(0, 0)
tcPr = cell._element.get_or_add_tcPr()
shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
tcPr.append(shd)

doc.save('scratch/test_ok.docx')
print('test script success!')
